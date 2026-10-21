import json, sys
from collections import OrderedDict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ENTRADA, SAIDA = sys.argv[1], sys.argv[2]
d = json.load(open(ENTRADA, encoding="utf-8"))
NUM = '#,##0.00;[Red]-#,##0.00'
PCT = '0.0%;[Red]-0.0%'
AZUL = PatternFill("solid", fgColor="D9E1F2")
CINZA = PatternFill("solid", fgColor="F2F2F2")
B = Font(bold=True)
fino = Side(style="thin", color="BFBFBF")

wb = Workbook()
ws = wb.active
ws.title = "DRE"
ws["A1"] = "NITAPLAST — DEMONSTRAÇÃO DO RESULTADO — SETEMBRO/2026 (antes do IRPJ/CSLL)"
ws["A1"].font = Font(bold=True, size=12)
ws.append([])
ws.append(["Linha", "AGO 26", "SET 26", "Variação R$", "Variação %", "% Rec. Bruta SET"])
for c in ws[3]:
    c.font = B; c.fill = AZUL; c.border = Border(bottom=fino)
principais = ("(+)", "(-)", "(=)", "LUCRO", "Resultado Operacional", "RESULTADO")
linha_rb = None
for desc, ago, sete in d["dre"]:
    if abs(ago) < 0.005 and abs(sete) < 0.005:
        continue
    r = ws.max_row + 1
    ws.cell(r, 1, desc.strip() if not desc.startswith("     ") else "   " + desc.strip())
    ws.cell(r, 2, round(ago, 2)).number_format = NUM
    ws.cell(r, 3, round(sete, 2)).number_format = NUM
    ws.cell(r, 4, f"=C{r}-B{r}").number_format = NUM
    ws.cell(r, 5, f'=IF(B{r}=0,"",(C{r}-B{r})/ABS(B{r}))').number_format = PCT
    if linha_rb is None:
        linha_rb = r
    ws.cell(r, 6, f"=C{r}/$C${linha_rb}").number_format = PCT
    if desc.strip().startswith(principais):
        for c in ws[r]:
            c.font = B
        if not desc.startswith("(+)Receita"):
            for c in ws[r]:
                c.fill = CINZA
notas = [
    "Notas:",
    "• CPV Matriz inclui estorno de R$ 109.923,88 da NF 94251 (Greatland), lançada em duplicidade no custo de agosto (erro de lançamento do sistema; agosto não foi alterado).",
    "• Resultado não operacional: estorno da provisão de receita de agosto (R$ 400.000,00).",
    "• Aluguel Victors: parcelas de junho a agosto (R$ 85.740,30) reconhecidas em setembro; aluguel rateado por CC conforme o Softdib.",
    "• Previdência Mongeral: particular do Marcos, mantida no resultado e evidenciada em linha própria.",
    "• Estoque Matriz pelo registro de inventário de 30/09/2026 (R$ 5.202.864,26); produto em elaboração igual a 31/08 (R$ 107.919,59), a confirmar.",
    "• Rendimento das aplicações Trust DI e Greencred sujeito aos extratos de setembro (ainda não recebidos).",
]
ws.append([])
for n in notas:
    ws.append([n])
    ws.cell(ws.max_row, 1).font = Font(bold=n == "Notas:", italic=n != "Notas:")
ws.column_dimensions["A"].width = 58
for col in "BCD":
    ws.column_dimensions[col].width = 16
ws.column_dimensions["E"].width = 12
ws.column_dimensions["F"].width = 16
ws.freeze_panes = "B4"

# DRE no layout da planilha de envio ("DRE Conferência — Nitaplast", aba do mês).
v0 = {desc.strip(): sete for desc, ago, sete in d["dre"]}
g = lambda k: round(v0.get(k, 0.0), 2)
e = wb.create_sheet("DRE Envio", 0)
e["C1"] = "NITAPLAST"; e["C1"].font = Font(bold=True, size=12)
e["A2"] = "DEMONSTRAÇÃO DO RESULTADO DE SETEMBRO 2026"; e["A2"].font = Font(bold=True, size=12)
e["C4"] = "COMPOSIÇÃO DRE"; e["C4"].font = B
cat = lambda k: g(k)
desp_fin_liq = round(g("Despesas Financeiras") - g("(-)Receitas Financeiras"), 2)
desp_oper = round(g("(-)Despesas Operacionais") - g("Outras receitas operacionais"), 2)
linhas_envio = [
    ("(+)Receita Operacional Bruta", g("(+)Receita Operacional Bruta"), True),
    ("     Receita Venda Produção Matriz", g("Receita Venda Produção Matriz"), False),
    ("     Receita Revenda Matriz", g("Receita Revenda Matriz"), False),
    ("     Receita Venda de Serviços", 0, False),
    ("     Receita Venda Produção Filial", g("Receita Venda Produção Filial"), False),
    ("     Receita Revenda Filial", g("Receita Revenda Filial"), False),
    None,
    ("(-)Deduções da Receita Bruta", g("(-)Deduções da Receita Bruta"), True),
    ("Devoluções -  Matriz", g("Devoluções - Matriz"), False),
    ("Devoluções - Filial", g("Devoluções - Filial"), False),
    ("IPI Matriz", g("IPI Matriz"), False), ("ICMS Matriz", g("ICMS Matriz"), False),
    ("Pis Matriz", g("Pis Matriz"), False), ("Cofins Matriz", g("Cofins Matriz"), False),
    ("ICMS ST", g("ICMS ST"), False), ("ICMS s/ vendas Filial", g("ICMS s/ vendas Filial"), False),
    ("IPI Filial", g("IPI Filial"), False), ("Pis Filial", g("Pis Filial"), False), ("Cofins Filial", g("Cofins Filial"), False),
    ("(-) Custo TOTAL", g("(-) Custo TOTAL"), True),
    ("(-) CPV Matriz", g("(-) CPV Matriz"), False), ("(-) CMV Matriz", 0, False),
    ("(-) CPV Filial", g("(-) CPV Filial"), False), ("(-) CMV Filial", 0, False),
    ("LUCRO BRUTO", g("LUCRO BRUTO"), True),
    None,
    ("(-)Despesas Operacionais", round(desp_oper + desp_fin_liq, 2), True),
    ("Despesas Administrativas", cat("Despesas Administrativas"), False),
    ("Despesas com Serviço - NPLog", cat("Despesas com Serviço - NPLog"), False),
    ("Despesas Comerciais", cat("Despesas Comerciais"), False),
    ("Despesas Produção", cat("Despesas Produção"), False),
    ("Despesas Veiculos", cat("Despesas Veiculos"), False),
    ("Despesas com Imobilizado", cat("Despesas com Imobilizado"), False),
    ("Despesas com Industrialização ", cat("Despesas com Industrialização"), False),
    ("Despesas comercial SP", cat("Despesas comercial SP"), False),
    ("Previdência particular — Marcos (Mongeral)", cat("Previdência particular — Marcos (Mongeral)"), False),
    ("(", desp_oper, False),
    ("Despesas Financeiras", g("Despesas Financeiras"), False),
    ("(-)Receitas Financeiras", g("(-)Receitas Financeiras"), False),
    ("Despesas Financeiras Líquidas", desp_fin_liq, True),
    None,
    ("TOTAL DAS DESPESAS OPERACIONAIS", round(desp_oper + desp_fin_liq, 2), True),
    None,
    ("Resultado Operacional", g("Resultado Operacional"), True),
    None,
    ("RESULTADO NÃO OPERACIONAL", g("RESULTADO NÃO OPERACIONAL"), True),
    ("       Receita de Alienação Imobilizado", 0, False),
    ("       Custo na Baixa/Alienação Imobilizado", 0, False),
    ("      Provisão para Custos - meses anteriores", 0, False),
    ("      Provisão para Custos - 07/2026", 0, False),
    ("     Reversão De provisão  Custos - 07/2026", 0, False),
    ("      Recuperação de Despesas Vinter", 0, False),
    ("Ganhos ou Perdas com Sinistros de Bens", 0, False),
    ("Outras Receitas não Operacionais (estorno da provisão de receita de 08/2026)", g("RESULTADO NÃO OPERACIONAL"), False),
    ("       Resultado da Equivalência Patrimonial", 0, False),
    None,
    ("LUCRO LIQUIDO (antes IRPJ/CSLL)", g("LUCRO/PREJUÍZO LÍQUIDO (antes IRPJ/CSLL)"), True),
]
outras_rec = g("Outras receitas operacionais")
r = 5
for item in linhas_envio:
    if item is None:
        r += 1; continue
    nome, valor, forte = item
    e.cell(r, 3, nome); e.cell(r, 4, valor).number_format = NUM
    if nome.startswith(("(+)Receita", "(-)Deduções", "(-) Custo", "LUCRO BRUTO", "(-)Despesas Oper", "Resultado Operacional", "LUCRO LIQUIDO")):
        e.cell(r, 5, f"=IF($D$5=0,0,D{r}/$D$5)").number_format = PCT
    if forte:
        e.cell(r, 3).font = B; e.cell(r, 4).font = B
    r += 1
if abs(outras_rec) > 0.005:
    e.cell(r + 1, 3, "Obs.: outras receitas operacionais (reembolsos) no resultado operacional"); e.cell(r + 1, 4, outras_rec).number_format = NUM
e.column_dimensions["C"].width = 62; e.column_dimensions["D"].width = 18; e.column_dimensions["E"].width = 10

# Análise horizontal e vertical (mesmo layout de "DRE - ANALISTE HORIZONTAL E VERTICAL 082026.xlsx").
v = {desc.strip(): (ago, sete) for desc, ago, sete in d["dre"]}
def par(nome):
    return v.get(nome, (0.0, 0.0))
fin = tuple(-(a - b) for a, b in zip(par("Despesas Financeiras"), par("(-)Receitas Financeiras")))
layout = [("(+) Receita Operacional Bruta", par("(+)Receita Operacional Bruta"), True)]
for k in ["Receita Venda Produção Matriz", "Receita Revenda Matriz", "Receita Venda Produção Filial", "Receita Revenda Filial"]:
    layout.append((k, par(k), False))
layout.append(("(-) Deduções da Receita Bruta", par("(-)Deduções da Receita Bruta"), True))
for k in ["Devoluções - Matriz", "Devoluções - Filial", "IPI Matriz", "ICMS Matriz", "Pis Matriz", "Cofins Matriz", "ICMS ST", "ICMS s/ vendas Filial", "IPI Filial", "Pis Filial", "Cofins Filial"]:
    layout.append((k, par(k), False))
layout += [("(=) Receita Operacional Líquida", par("(=) Receita Operacional Líquida"), True),
           ("(-) CPV / CMV Total", par("(-) Custo TOTAL"), True),
           ("CPV Matriz", par("(-) CPV Matriz"), False), ("CPV Filial", par("(-) CPV Filial"), False),
           ("(=) Lucro Bruto", par("LUCRO BRUTO"), True),
           ("(-) Despesas Operacionais", par("(-)Despesas Operacionais"), True)]
for k in ["Despesas Administrativas", "Despesas com Serviço - NPLog", "Despesas Comerciais", "Despesas Produção", "Despesas Veiculos", "Despesas Barracão", "Despesas com Imobilizado", "Despesas com Industrialização", "Despesas Tributárias", "Despesas comercial SP", "Previdência particular — Marcos (Mongeral)", "Outras despesas operacionais sem classificação gerencial"]:
    if k in v and (abs(v[k][0]) > 0.005 or abs(v[k][1]) > 0.005):
        layout.append((k, par(k), False))
layout += [("Resultado Financeiro Líquido", fin, True), ("Despesas Financeiras", par("Despesas Financeiras"), False),
           ("(-) Receitas Financeiras", par("(-)Receitas Financeiras"), False)]
if any(abs(x) > 0.005 for x in par("Outras receitas operacionais")):
    layout.append(("Outras receitas operacionais", par("Outras receitas operacionais"), False))
layout += [("(=) Resultado Operacional", par("Resultado Operacional"), True),
           ("Resultado não operacional", par("RESULTADO NÃO OPERACIONAL"), True),
           ("(=) Lucro / Prejuízo Líquido (antes IRPJ/CSLL)", par("LUCRO/PREJUÍZO LÍQUIDO (antes IRPJ/CSLL)"), True)]
h = wb.create_sheet("Análise H e V", 1)
h.append(["Composição DRE", "AGO 26", "AV 08", "SET 26", "AV 09", "Diferença R$", "Variação %", "Variação p.p."])
for c in h[1]:
    c.font = B; c.fill = AZUL; c.border = Border(bottom=fino); c.alignment = Alignment(horizontal="center")
for nome, (a, s9), forte in layout:
    r = h.max_row + 1
    h.cell(r, 1, nome if forte else "   " + nome)
    h.cell(r, 2, round(a, 2)).number_format = NUM
    h.cell(r, 3, f"=IF($B$2=0,0,B{r}/$B$2)").number_format = PCT
    h.cell(r, 4, round(s9, 2)).number_format = NUM
    h.cell(r, 5, f"=IF($D$2=0,0,D{r}/$D$2)").number_format = PCT
    h.cell(r, 6, f"=D{r}-B{r}").number_format = NUM
    h.cell(r, 7, f'=IF(B{r}=0,"",(D{r}-B{r})/ABS(B{r}))').number_format = PCT
    h.cell(r, 8, f"=(E{r}-C{r})*100").number_format = '0.00;[Red]-0.00'
    if forte:
        for c in h[r]:
            c.font = B; c.fill = CINZA
h.column_dimensions["A"].width = 52
for col in "BDF":
    h.column_dimensions[col].width = 16
for col in "CEGH":
    h.column_dimensions[col].width = 12
h.freeze_panes = "B2"

# Composição lançamento a lançamento, uma aba por linha da DRE (mesmo modelo de agosto).
grupos = OrderedDict()
for x in d["itens"]:
    grupos.setdefault(x["grupo"], OrderedDict()).setdefault(x["conta"], []).append(x)
ORDEM = ["CPV Matriz", "CPV Filial", "Despesas Administrativas", "Despesas com Serviço - NPLog", "Despesas Comerciais", "Despesas Produção", "Despesas Veiculos", "Despesas com Imobilizado", "Despesas com Industrialização", "Despesas Tributárias", "Despesas comercial SP", "Previdência particular — Marcos (Mongeral)", "Outras despesas operacionais sem classificação gerencial", "Despesas Financeiras", "Receitas Financeiras", "Outras Receitas Operacionais", "Resultado Não Operacional"]
ABA = {"Outras despesas operacionais sem classificação gerencial": "Desp. sem classificação", "Previdência particular — Marcos (Mongeral)": "Previdência Marcos"}
for grupo, contas in sorted(grupos.items(), key=lambda kv: ORDEM.index(kv[0]) if kv[0] in ORDEM else 99):
    t = wb.create_sheet(ABA.get(grupo, grupo)[:31])
    t["A1"] = f"{grupo.upper()} — SETEMBRO/2026"
    t["A1"].font = Font(bold=True, size=12)
    t.append(["Data", "CC", "Histórico", "Documento", "R$"])
    for c in t[2]:
        c.font = B; c.border = Border(bottom=fino)
    subtotais = []
    ordem = sorted(contas.items(), key=lambda kv: -abs(sum(i["valor"] for i in kv[1])))
    for conta, itens in ordem:
        t.append([f"{conta} {itens[0]['nomeConta']}".strip()])
        r0 = t.max_row
        for c in t[r0]:
            c.font = B
        t.cell(r0, 1).fill = AZUL
        for i in sorted(itens, key=lambda i: -abs(i["valor"])):
            t.append([i["data"], i["cc"], i["historico"], i["documento"], i["valor"]])
            t.cell(t.max_row, 5).number_format = NUM
        r = t.max_row + 1
        t.cell(r, 3, f"Subtotal {conta}").font = B
        t.cell(r, 5, f"=SUM(E{r0 + 1}:E{r - 1})").number_format = NUM
        t.cell(r, 5).font = B
        subtotais.append(f"E{r}")
    t.append([])
    r = t.max_row + 1
    t.cell(r, 3, f"TOTAL {grupo.upper()}").font = Font(bold=True, size=11)
    t.cell(r, 5, "=" + "+".join(subtotais)).number_format = NUM
    t.cell(r, 5).font = B
    t.cell(r, 3).fill = CINZA; t.cell(r, 5).fill = CINZA
    for col, w in zip("ABCDE", (12, 8, 70, 45, 16)):
        t.column_dimensions[col].width = w
    t.freeze_panes = "A3"
    for row in t.iter_rows(min_row=3):
        row[2].alignment = Alignment(wrap_text=False)

wb.save(SAIDA)
print("ok", SAIDA, [s.title for s in wb.worksheets])
